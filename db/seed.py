"""
Seed script: reads the Creative Quality Evals Criteria Excel file and the
generations CSV, then populates a SQLite database at db/evals.db.
Run this script to recreate the database from scratch.

Usage:
    python db/seed.py
"""

import csv
import json
import re
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

EXCEL_PATH = Path(__file__).parent.parent / "data" / "Creative Quality Evals Criteria_in progress [WIP].xlsx"
GENERATIONS_CSV = Path(__file__).parent.parent / "data" / "generations_output.csv"
DB_PATH = Path(__file__).parent / "evals.db"

NOW = datetime.now(timezone.utc).isoformat()

# ---------------------------------------------------------------------------
# Sheet configurations
# Each sheet has a fixed column layout; see COLUMN GUIDE below.
#
# col_eval_def: 0-based column index of the criteria-level definition text.
# After that column the eval scoring data repeats in blocks of 4:
#   [title, definition, example1, example2] per score level (1-4)
# For Y/N criteria the same columns are used: block 1 = Yes, block 2 = No.
# ---------------------------------------------------------------------------
SHEET_CONFIGS = [
    {
        "name": "Amazon Title",
        "context": "Marketplace", "marketplace_tag": "Amazon", "content_type": "Title",
        "header_row": 3, "data_start": 4,
        "col_category": 0, "col_weight": 1, "col_criteria": 3,
        "col_qtype": 6, "col_eval_def": 9,
    },
    {
        "name": "Amazon Description",
        "context": "Marketplace", "marketplace_tag": "Amazon", "content_type": "Description",
        "header_row": 3, "data_start": 4,
        "col_category": 0, "col_weight": 1, "col_criteria": 3,
        "col_qtype": 6, "col_eval_def": 9,
    },
    {
        "name": "Universal Description",
        "context": "Universal", "content_type": "Description",
        "header_row": 3, "data_start": 4,
        "col_category": 0, "col_weight": 1, "col_criteria": 2,
        "col_qtype": 4, "col_eval_def": 7,
    },
    {
        "name": "Universal Title",
        "context": "Universal", "content_type": "Title",
        "header_row": 4, "data_start": 5,
        "col_category": 0, "col_weight": 1, "col_criteria": 2,
        "col_qtype": 5, "col_eval_def": 8,
    },
    {
        "name": "Universal Bullets Specs",
        "context": "Universal", "content_type": "Bullets/Specs",
        "header_row": 3, "data_start": 4,
        "col_category": 0, "col_weight": 1, "col_criteria": 2,
        "col_qtype": 5, "col_eval_def": 8,
    },
    {
        "name": "Universal Meta Description",
        "context": "Universal", "content_type": "Meta Description",
        "header_row": 3, "data_start": 4,
        "col_category": 0, "col_weight": 1, "col_criteria": 2,
        "col_qtype": 5, "col_eval_def": 7,
    },
]

QTYPE_MAP = {
    "y/n": "yes-no",
    "yn": "yes-no",
    "yes/no": "yes-no",
    "scale": "numerical-scale",
    "how many": "numerical-count",
}


def slugify(text: str) -> str:
    text = text.lower().strip()
    text = re.sub(r"[^\w\s-]", "", text)
    text = re.sub(r"[\s_]+", "-", text)
    text = re.sub(r"-+", "-", text)
    return text.strip("-")[:80]


def clean(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    # Replace common Windows-1252 replacement chars used in the Excel file
    s = s.replace("�", "'").replace("’", "'").replace("‘", "'")
    s = s.replace("“", '"').replace("”", '"')
    s = s.replace("–", "-").replace("—", "-")
    s = s.replace(" ", " ")  # line separator
    return s if s else None


def parse_eval_definition(row: list, qtype: str, eval_col: int) -> dict:
    """Parse the eval definition columns into the appropriate typed structure."""

    def g(offset: int):
        idx = eval_col + offset
        return clean(row[idx]) if idx < len(row) else None

    if qtype == "yes-no":
        # Block 1 = Yes (offsets 1-4), Block 2 = No (offsets 5-8)
        yes_title = g(1)
        yes_def = g(2)
        yes_ex1 = g(3)
        yes_ex2 = g(4)
        no_title = g(5)
        no_def = g(6)
        no_ex1 = g(7)
        no_ex2 = g(8)

        result: dict = {}
        if yes_def:
            result["definition_yes"] = yes_def
        if no_def:
            result["definition_no"] = no_def

        yes_examples = [e for e in [yes_ex1, yes_ex2] if e]
        if yes_examples:
            result["yes_examples"] = yes_examples
        no_examples = [e for e in [no_ex1, no_ex2] if e]
        if no_examples:
            result["no_examples"] = no_examples

        return result

    elif qtype == "numerical-scale":
        result = {}
        for score in range(1, 5):
            base = (score - 1) * 4
            title = g(1 + base)
            definition = g(2 + base)
            ex1 = g(3 + base)
            ex2 = g(4 + base)
            if not title and not definition:
                continue
            entry: dict = {}
            if title:
                entry["title"] = title
            if definition:
                entry["definition"] = definition
            if ex1:
                entry["example_1"] = ex1
            if ex2:
                entry["example_2"] = ex2
            result[f"score_{score}"] = entry
        return result

    elif qtype == "numerical-count":
        # Buckets are 0, 1, 2, 3+
        buckets = ["0", "1", "2", "3+"]
        bucket_definitions = {}
        for i, bucket in enumerate(buckets):
            val = g(1 + i)
            if val:
                bucket_definitions[bucket] = val
        return {
            "buckets": buckets,
            "bucket_definitions": bucket_definitions,
        }

    return {}


def parse_sheet(ws, config: dict) -> list[dict]:
    rows_data = list(ws.iter_rows(min_row=config["data_start"], values_only=True))

    results = []
    last_category = None
    last_weight = None
    seen_ids: set[str] = set()

    for raw_row in rows_data:
        row = list(raw_row)

        # Pad row to at least 30 cols so index access is safe
        while len(row) < 30:
            row.append(None)

        criteria_name = clean(row[config["col_criteria"]])
        if not criteria_name:
            continue

        qtype_raw = clean(row[config["col_qtype"]])
        if not qtype_raw:
            continue
        qtype = QTYPE_MAP.get(qtype_raw.lower().strip(), None)
        if not qtype:
            continue

        category = clean(row[config["col_category"]]) or last_category
        last_category = category

        raw_weight = row[config["col_weight"]]
        if raw_weight is not None and str(raw_weight).strip():
            try:
                last_weight = float(raw_weight)
            except (ValueError, TypeError):
                pass
        weight = last_weight if last_weight is not None else 1.0

        eval_col = config["col_eval_def"]
        criteria_def = clean(row[eval_col])

        eval_definition = parse_eval_definition(row, qtype, eval_col)

        # Build a unique ID from criteria name + content type
        base_id = slugify(criteria_name)
        content_slug = slugify(config["content_type"])
        candidate_id = f"{base_id}-{content_slug}"
        # Deduplicate within this seed run
        uid = candidate_id
        counter = 2
        while uid in seen_ids:
            uid = f"{candidate_id}-{counter}"
            counter += 1
        seen_ids.add(uid)

        record: dict = {
            "id": uid,
            "context": config["context"],
            "content_type": config["content_type"],
            "criteria_category": category or "Uncategorized",
            "criteria_name": criteria_name,
            "criteria_definition": criteria_def or "",
            "criteria_type": qtype,
            "eval_definition": eval_definition,
            "weight": weight,
            "active": True,
            "marketplace_tag": config.get("marketplace_tag"),
            "brand_tag": config.get("brand_tag"),
            "industry_tag": config.get("industry_tag"),
            "created_at": NOW,
            "updated_at": NOW,
        }
        results.append(record)

    return results


def create_schema(conn: sqlite3.Connection) -> None:
    conn.execute("""
        CREATE TABLE IF NOT EXISTS criteria (
            id              TEXT PRIMARY KEY,
            context         TEXT NOT NULL,
            content_type    TEXT NOT NULL,
            criteria_category TEXT,
            criteria_name   TEXT NOT NULL,
            criteria_definition TEXT,
            criteria_type   TEXT NOT NULL,
            eval_definition TEXT NOT NULL,
            weight          REAL NOT NULL DEFAULT 1.0,
            active          INTEGER NOT NULL DEFAULT 1,
            marketplace_tag TEXT,
            brand_tag       TEXT,
            industry_tag    TEXT,
            customer        TEXT,
            brand           TEXT,
            custom_tags     TEXT,
            created_at      TEXT NOT NULL,
            updated_at      TEXT NOT NULL
        )
    """)
    conn.commit()


def insert_criteria(conn: sqlite3.Connection, records: list[dict]) -> None:
    for r in records:
        conn.execute(
            """
            INSERT OR REPLACE INTO criteria
              (id, context, content_type, criteria_category, criteria_name,
               criteria_definition, criteria_type, eval_definition,
               weight, active, marketplace_tag, brand_tag, industry_tag,
               customer, brand, custom_tags, created_at, updated_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                r["id"],
                r["context"],
                r["content_type"],
                r["criteria_category"],
                r["criteria_name"],
                r["criteria_definition"],
                r["criteria_type"],
                json.dumps(r["eval_definition"]),
                r["weight"],
                1 if r["active"] else 0,
                r.get("marketplace_tag"),
                r.get("brand_tag"),
                r.get("industry_tag"),
                r.get("customer"),
                r.get("brand"),
                json.dumps(r.get("custom_tags") or {}),
                r["created_at"],
                r["updated_at"],
            ),
        )
    conn.commit()


def create_eval_results_schema(conn: sqlite3.Connection) -> None:
    conn.execute("""
        CREATE TABLE IF NOT EXISTS eval_results (
            id              TEXT PRIMARY KEY,
            generation_id   TEXT NOT NULL,
            criterion_id    TEXT NOT NULL,
            criterion_name  TEXT,
            desired_score   TEXT,
            score           TEXT,
            rationale       TEXT,
            evidence        TEXT,
            product_name    TEXT,
            run_at          TEXT NOT NULL
        )
    """)
    conn.commit()


def create_generations_schema(conn: sqlite3.Connection) -> None:
    conn.execute("""
        CREATE TABLE IF NOT EXISTS generations (
            generation_id       TEXT PRIMARY KEY,
            model               TEXT,
            created_at          INTEGER,
            system_prompt       TEXT,
            last_user_message   TEXT,
            few_shot_count      INTEGER,
            temperature         REAL,
            max_tokens          INTEGER,
            response_content    TEXT,
            prompt_tokens       INTEGER,
            completion_tokens   INTEGER,
            total_tokens        INTEGER,
            finish_reason       TEXT,
            is_valid            INTEGER,
            req_json            TEXT,
            resp_json           TEXT
        )
    """)
    conn.commit()


def _int(val: str) -> int | None:
    try:
        return int(val) if val else None
    except (ValueError, TypeError):
        return None


def _float(val: str) -> float | None:
    try:
        return float(val) if val else None
    except (ValueError, TypeError):
        return None


def seed_generations(conn: sqlite3.Connection) -> int:
    if not GENERATIONS_CSV.exists():
        print(f"  WARNING: generations CSV not found at {GENERATIONS_CSV}, skipping")
        return 0

    with open(GENERATIONS_CSV, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    count = 0
    for row in rows:
        is_valid_raw = row.get("resp_is_valid", "").strip().lower()
        is_valid = 1 if is_valid_raw == "true" else (0 if is_valid_raw == "false" else None)

        conn.execute(
            """
            INSERT OR REPLACE INTO generations
              (generation_id, model, created_at, system_prompt, last_user_message,
               few_shot_count, temperature, max_tokens, response_content,
               prompt_tokens, completion_tokens, total_tokens,
               finish_reason, is_valid, req_json, resp_json)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                row.get("generation_id"),
                row.get("resp_model") or None,
                _int(row.get("resp_created")),
                row.get("req_system_prompt") or None,
                row.get("req_last_user_message") or None,
                _int(row.get("req_few_shot_count")),
                _float(row.get("req_temperature")),
                _int(row.get("req_max_tokens")),
                row.get("resp_content") or None,
                _int(row.get("resp_prompt_tokens")),
                _int(row.get("resp_completion_tokens")),
                _int(row.get("resp_total_tokens")),
                row.get("resp_finish_reason") or None,
                is_valid,
                row.get("req_json") or None,
                row.get("resp_json") or None,
            ),
        )
        count += 1
    conn.commit()
    return count


def seed_type_mapping(conn: sqlite3.Connection) -> None:
    """
    Maps criteria content_type → generation inferred type.
    One criteria content_type can map to multiple generation types.
    """
    conn.execute("""
        CREATE TABLE IF NOT EXISTS type_mapping (
            criteria_content_type   TEXT NOT NULL,
            generation_type         TEXT NOT NULL,
            PRIMARY KEY (criteria_content_type, generation_type)
        )
    """)
    defaults = [
        ("Title",          "Title"),
        ("Description",    "Description"),
        ("Bullets/Specs",  "Bullets"),
        ("Meta Description", "Description"),
    ]
    conn.executemany(
        "INSERT OR IGNORE INTO type_mapping (criteria_content_type, generation_type) VALUES (?,?)",
        defaults,
    )
    conn.commit()


def main() -> None:
    if not EXCEL_PATH.exists():
        print(f"Excel file not found: {EXCEL_PATH}", file=sys.stderr)
        sys.exit(1)

    print(f"Reading: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(str(EXCEL_PATH))

    all_records: list[dict] = []
    for config in SHEET_CONFIGS:
        sheet_name = config["name"]
        if sheet_name not in wb.sheetnames:
            print(f"  WARNING: sheet '{sheet_name}' not found, skipping")
            continue
        ws = wb[sheet_name]
        records = parse_sheet(ws, config)
        print(f"  {sheet_name}: {len(records)} criteria")
        all_records.extend(records)

    print(f"\nTotal criteria parsed: {len(all_records)}")

    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(DB_PATH))

    # Drop and recreate for a clean seed
    conn.execute("DROP TABLE IF EXISTS criteria")
    conn.execute("DROP TABLE IF EXISTS generations")
    conn.execute("DROP TABLE IF EXISTS eval_results")
    conn.execute("DROP TABLE IF EXISTS type_mapping")
    create_schema(conn)
    insert_criteria(conn, all_records)

    create_generations_schema(conn)
    gen_count = seed_generations(conn)
    create_eval_results_schema(conn)
    print(f"Generations seeded: {gen_count} rows")

    seed_type_mapping(conn)
    print("Type mapping seeded")

    conn.close()
    print(f"Database written: {DB_PATH}")


if __name__ == "__main__":
    main()
